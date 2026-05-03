"""Extract highlight (shading) stats from expert LLM grading .docx files.

Word stores colored highlights as ``w:shd`` with ``w:fill`` hex RGB (not ``w:highlight``).
Legend in doc: Red / Yellow / Green (see rubric 0--5 separate from highlight semantics).

Run:

    python -m viz._analyze_llm_grading_docx_highlights
"""

from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
WNS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# Observed fills in repo docs (uppercase hex, no #)
FILL_WHITE = "FFFFFF"
FILLS = {
    "FF0000": "red",
    "FFFF00": "yellow",
    "00FF00": "green",
}


def enrich_counts(counts: dict[str, int]) -> dict[str, float | int]:
    tot = sum(counts.values())
    rgb = counts.get("red", 0) + counts.get("yellow", 0) + counts.get("green", 0)
    return {
        "total_chars": tot,
        "highlighted_rgb_chars": rgb,
        "pct_highlighted_of_reasoning": round(100.0 * rgb / tot, 2) if tot else 0.0,
    }


def _run_fill_hex(run) -> str | None:
    el = run._element
    rpr = el.find(WNS + "rPr")
    if rpr is None:
        return None
    shd = rpr.find(WNS + "shd")
    if shd is None:
        return None
    fill = shd.get(WNS + "fill") or shd.get("fill")
    return fill.upper() if fill else None


def _classify_paragraph(text: str) -> str:
    s = text.strip()
    if not s:
        return "blank"
    if s.startswith('"leading_interpretation'):
        return "q2"
    if s.startswith('"alternative_analysis'):
        return "q3"
    if s.startswith("LLM's Classification"):
        return "meta"
    if s.startswith("Correct Class"):
        return "meta"
    if s.startswith("Your Grading"):
        return "meta"
    if s.startswith("Note:"):
        return "meta"
    return "intro_other"


def _count_runs_fills_in_paragraph(para) -> dict[str, int]:
    """Tally character counts by highlight category for one paragraph."""
    out: dict[str, int] = defaultdict(int)
    for run in para.runs:
        t = run.text
        if not t:
            continue
        fill = _run_fill_hex(run)
        if fill in (None, FILL_WHITE):
            cat = "unhighlighted"
        elif fill in FILLS:
            cat = FILLS[fill]
        else:
            cat = f"other_fill_{fill}"
        out[cat] += len(t)
    return dict(out)


def _merge_fill_counts(a: dict[str, int], b: dict[str, int]) -> dict[str, int]:
    out: dict[str, int] = defaultdict(int)
    for d in (a, b):
        for k, v in d.items():
            out[k] += v
    return dict(out)


def _parse_colon_field(text: str, prefix: str) -> str:
    s = text.strip()
    if not s.lower().startswith(prefix.lower()):
        return ""
    rest = s.split(":", 1)
    if len(rest) < 2:
        return ""
    return rest[1].strip()


def _labels_equivalent(a: str, b: str) -> bool:
    return " ".join(a.lower().split()) == " ".join(b.lower().split())


def per_model_merged_fills_in_doc_order(path: Path) -> list[dict[str, int]]:
    """Return one Q2+Q3 fill tally per model block, in document order (typ. 13 blocks)."""
    d = Document(path)
    q2_c: dict[str, int] | None = None
    q3_c: dict[str, int] | None = None
    pred: str | None = None
    blocks: list[dict[str, int]] = []
    for para in d.paragraphs:
        t = para.text
        s = t.strip()
        if not s:
            continue
        if s.startswith('"leading_interpretation'):
            q2_c = _count_runs_fills_in_paragraph(para)
            continue
        if s.startswith('"alternative_analysis'):
            q3_c = _count_runs_fills_in_paragraph(para)
            continue
        if s.startswith("LLM's Classification"):
            pred = _parse_colon_field(t, "LLM's Classification")
            continue
        if s.startswith("Correct Class"):
            if q2_c is not None and q3_c is not None:
                blocks.append(_merge_fill_counts(q2_c, q3_c))
            q2_c = q3_c = None
            pred = None
            continue
    return blocks


def _para_runs_to_color_segments(para) -> list[tuple[str, str]]:
    """Adjacent runs with the same fill merged; categories match ``FILLS`` + unhighlighted."""
    segments: list[tuple[str, str]] = []
    for run in para.runs:
        text = run.text
        if not text:
            continue
        fill = _run_fill_hex(run)
        if fill in (None, FILL_WHITE):
            cat = "unhighlighted"
        elif fill in FILLS:
            cat = FILLS[fill]
        else:
            cat = "unhighlighted"
        if segments and segments[-1][1] == cat:
            segments[-1] = (segments[-1][0] + text, cat)
        else:
            segments.append((text, cat))
    return segments


def _trim_segment_prefix(segments: list[tuple[str, str]], prefix: str) -> list[tuple[str, str]]:
    """Drop ``prefix`` from the start of concatenated segments (Word copies JSON key lines)."""
    if not prefix or not segments:
        return segments
    full = "".join(t for t, _ in segments)
    if not full.startswith(prefix):
        return segments
    rem = len(prefix)
    out: list[tuple[str, str]] = []
    for text, cat in segments:
        if rem <= 0:
            out.append((text, cat))
            continue
        if len(text) <= rem:
            rem -= len(text)
            continue
        out.append((text[rem:], cat))
        rem = 0
    return out


def _trim_segment_suffix_quote(segments: list[tuple[str, str]]) -> list[tuple[str, str]]:
    """Remove a single trailing ASCII ``"`` when Word closes the JSON string in the paragraph."""
    if not segments:
        return segments
    text, cat = segments[-1]
    if text.endswith('"') and len(text) >= 1:
        segments = segments[:-1] + [(text[:-1], cat)]
    return segments


def per_model_reasoning_color_segments_in_doc_order(
    path: Path,
) -> list[dict[str, list[tuple[str, str]]]]:
    """Per model block: colored segments for Part B Q2/Q3 only (doc paragraph order)."""
    d = Document(path)
    blocks: list[dict[str, list[tuple[str, str]]]] = []
    q2_seg: list[tuple[str, str]] | None = None
    q3_seg: list[tuple[str, str]] | None = None
    for para in d.paragraphs:
        t = para.text
        s = t.strip()
        if not s:
            continue
        if s.startswith('"leading_interpretation'):
            q2_seg = _para_runs_to_color_segments(para)
            continue
        if s.startswith('"alternative_analysis'):
            q3_seg = _para_runs_to_color_segments(para)
            continue
        if s.startswith("Correct Class"):
            if q2_seg is not None and q3_seg is not None:
                q2_trim = _trim_segment_suffix_quote(
                    _trim_segment_prefix(
                        q2_seg, '"leading_interpretation_and_support": "'
                    )
                )
                q3_trim = _trim_segment_suffix_quote(
                    _trim_segment_prefix(q3_seg, '"alternative_analysis": "')
                )
                blocks.append(
                    {
                        "leading_interpretation_and_support": q2_trim,
                        "alternative_analysis": q3_trim,
                    }
                )
            q2_seg = q3_seg = None
            continue
    return blocks


def dominant_reasoning_fill(merged: dict[str, int]) -> str:
    """Highlight bucket with the most characters in Q2+Q3 (ties favor green, then yellow, red, unhighlighted)."""
    keys = ("green", "yellow", "red", "unhighlighted")
    counts = {k: merged.get(k, 0) for k in keys}
    for k, v in merged.items():
        if k.startswith("other_fill_"):
            counts["unhighlighted"] += v
    tie_break = {"green": 3, "yellow": 2, "red": 1, "unhighlighted": 0}
    return max(keys, key=lambda k: (counts[k], tie_break[k]))


def analyze_docx_by_correctness(path: Path) -> dict:
    """Per model block, merge Q2+Q3 fill counts; split by (LLM pred == gold)."""
    d = Document(path)
    correct: dict[str, int] = defaultdict(int)
    incorrect: dict[str, int] = defaultdict(int)
    n_correct = 0
    n_incorrect = 0
    n_unresolved = 0
    q2_c: dict[str, int] | None = None
    q3_c: dict[str, int] | None = None
    pred: str | None = None

    for para in d.paragraphs:
        t = para.text
        s = t.strip()
        if not s:
            continue
        if s.startswith('"leading_interpretation'):
            q2_c = _count_runs_fills_in_paragraph(para)
            continue
        if s.startswith('"alternative_analysis'):
            q3_c = _count_runs_fills_in_paragraph(para)
            continue
        if s.startswith("LLM's Classification"):
            pred = _parse_colon_field(t, "LLM's Classification")
            continue
        if s.startswith("Correct Class"):
            gold = _parse_colon_field(t, "Correct Class")
            if q2_c is None or q3_c is None:
                n_unresolved += 1
                pred = None
                continue
            merged = _merge_fill_counts(q2_c, q3_c)
            ok = pred is not None and gold != "" and _labels_equivalent(pred, gold)
            tgt = correct if ok else incorrect
            for k, v in merged.items():
                tgt[k] += v
            if ok:
                n_correct += 1
            else:
                n_incorrect += 1
            q2_c = q3_c = None
            pred = None
            continue

    both_c = dict(correct)
    both_i = dict(incorrect)

    def pct_histogram(counts: dict[str, int]) -> dict[str, float]:
        tot = sum(counts.values())
        if tot == 0:
            return {}
        return {k: round(100.0 * v / tot, 2) for k, v in sorted(counts.items(), key=lambda x: -x[1])}

    n_blocks = n_correct + n_incorrect
    return {
        "path": str(path.name),
        "n_blocks_classified": n_blocks,
        "n_correct_prediction": n_correct,
        "n_incorrect_prediction": n_incorrect,
        "fraction_blocks_correct": round(n_correct / n_blocks, 4) if n_blocks else 0.0,
        "n_unresolved_blocks": n_unresolved,
        "correct_chars_by_fill": both_c,
        "incorrect_chars_by_fill": both_i,
        "pct_correct": pct_histogram(both_c),
        "pct_incorrect": pct_histogram(both_i),
        "summary_correct": enrich_counts(both_c),
        "summary_incorrect": enrich_counts(both_i),
        "avg_chars_per_block_correct": (
            round(sum(both_c.values()) / n_correct, 1) if n_correct else None
        ),
        "avg_chars_per_block_incorrect": (
            round(sum(both_i.values()) / n_incorrect, 1) if n_incorrect else None
        ),
    }


def analyze_docx(path: Path) -> dict:
    d = Document(path)
    # chars per fill bucket and section (q2 = leading interp., q3 = alternative)
    per_section: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))

    for para in d.paragraphs:
        bucket = _classify_paragraph(para.text)
        if bucket not in ("q2", "q3"):
            continue
        for run in para.runs:
            t = run.text
            if not t:
                continue
            fill = _run_fill_hex(run)
            if fill in (None, FILL_WHITE):
                cat = "unhighlighted"
            elif fill in FILLS:
                cat = FILLS[fill]
            else:
                cat = f"other_fill_{fill}"

            per_section[bucket][cat] += len(t)

    def rollup(sec_keys: tuple[str, ...]) -> dict[str, int]:
        out: dict[str, int] = defaultdict(int)
        for sk in sec_keys:
            for k, v in per_section[sk].items():
                out[k] += v
        return dict(out)

    q2 = dict(per_section["q2"])
    q3 = dict(per_section["q3"])
    both = rollup(("q2", "q3"))

    def pct_histogram(counts: dict[str, int]) -> dict[str, float]:
        tot = sum(counts.values())
        if tot == 0:
            return {}
        return {k: round(100.0 * v / tot, 2) for k, v in sorted(counts.items(), key=lambda x: -x[1])}

    def enrich(counts: dict[str, int]) -> dict[str, float | int]:
        return enrich_counts(counts)

    n_models = sum(
        1
        for p in d.paragraphs
        if p.text.strip().startswith('"leading_interpretation')
    )

    return {
        "path": str(path.name),
        "n_models": n_models,
        "q2_leading_interpretation_chars_by_fill": q2,
        "q3_alternative_analysis_chars_by_fill": q3,
        "reasoning_both_parts_chars_by_fill": both,
        "pct_q2": pct_histogram(q2),
        "pct_q3": pct_histogram(q3),
        "pct_both": pct_histogram(both),
        "summary_q2": enrich(q2),
        "summary_q3": enrich(q3),
        "summary_both": enrich(both),
        "avg_chars_per_model_both": (
            round(sum(both.values()) / n_models, 1) if n_models else None
        ),
        "by_llm_correct_vs_gold": analyze_docx_by_correctness(path),
    }


def main() -> None:
    files = [
        ROOT / "LLM Answer Grading ZTF19abfqvbg.docx",
        ROOT / "LLM Answer Grading ZTF26aargnnp.docx",
    ]
    results = []
    for f in files:
        if not f.is_file():
            print("missing", f)
            continue
        results.append(analyze_docx(f))
    combined_both: dict[str, int] = defaultdict(int)
    for r in results:
        for k, v in r["reasoning_both_parts_chars_by_fill"].items():
            combined_both[k] += v

    comb_correct: dict[str, int] = defaultdict(int)
    comb_incorrect: dict[str, int] = defaultdict(int)
    n_corr = n_incorr = 0
    for r in results:
        bc = r.get("by_llm_correct_vs_gold") or {}
        n_corr += int(bc.get("n_correct_prediction", 0))
        n_incorr += int(bc.get("n_incorrect_prediction", 0))
        for k, v in (bc.get("correct_chars_by_fill") or {}).items():
            comb_correct[k] += v
        for k, v in (bc.get("incorrect_chars_by_fill") or {}).items():
            comb_incorrect[k] += v

    def pct(counts: dict[str, int]) -> dict[str, float]:
        tot = sum(counts.values())
        if tot == 0:
            return {}
        return {k: round(100.0 * v / tot, 2) for k, v in sorted(counts.items(), key=lambda x: -x[1])}

    combined_pct = pct(dict(combined_both))

    out = {
        "per_datapoint": results,
        "combined_two_datapoints": {
            "reasoning_both_parts_chars_by_fill": dict(combined_both),
            "pct_both": combined_pct,
            "total_reasoning_chars": sum(combined_both.values()),
            "summary_both": enrich_counts(dict(combined_both)),
            "n_model_blocks_counted": sum(r["n_models"] for r in results),
            "by_llm_correct_vs_gold_pooled": {
                "n_blocks_correct": n_corr,
                "n_blocks_incorrect": n_incorr,
                "correct_chars_by_fill": dict(comb_correct),
                "incorrect_chars_by_fill": dict(comb_incorrect),
                "pct_correct": pct(dict(comb_correct)),
                "pct_incorrect": pct(dict(comb_incorrect)),
                "summary_correct": enrich_counts(dict(comb_correct)),
                "summary_incorrect": enrich_counts(dict(comb_incorrect)),
            },
        },
        "legend_note": (
            "Highlights stored as w:shd fill in OOXML. Mapped: FF0000=red, FFFF00=yellow, "
            "00FF00=green; FFFFFF counted as unhighlighted within reasoning paragraphs."
        ),
    }
    print(json.dumps(out, indent=2))


if __name__ == "__main__":
    main()
